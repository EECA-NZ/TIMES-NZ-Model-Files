"""Tests for the VEDA Excel cell-type conversion rules."""

import pandas as pd
import pytest
from openpyxl import load_workbook
from prepare_times_nz.utilities import excel_writers
from prepare_times_nz.utilities.excel_writers import coerce_excel_data_value


def test_coerce_excel_data_value_preserves_integer_form_values():
    """Integer-form input is written as an integer, not as a float."""
    assert coerce_excel_data_value("2023") == 2023
    assert isinstance(coerce_excel_data_value("2023"), int)
    assert coerce_excel_data_value(-1) == -1


def test_coerce_excel_data_value_converts_decimal_and_scientific_values():
    """Decimal-form input is numeric and is limited to 15 significant figures."""
    assert coerce_excel_data_value("0.95") == 0.95
    assert isinstance(coerce_excel_data_value("0.95"), float)
    assert coerce_excel_data_value("1e-4") == 0.0001
    assert coerce_excel_data_value("2078.7684537684536") == 2078.76845376845


def test_coerce_excel_data_value_preserves_non_numeric_text():
    """Names, patterns, blanks, and non-finite numeric text stay as text."""
    assert coerce_excel_data_value("ELC_*") == "ELC_*"
    assert coerce_excel_data_value("") == ""
    assert coerce_excel_data_value("NaN") == "NaN"


@pytest.mark.parametrize("value", [None, float("nan"), pd.NA])
def test_coerce_excel_data_value_writes_missing_values_as_blank_cells(value):
    """Missing-data sentinels must not be converted to the string ``nan``."""
    assert coerce_excel_data_value(value) is None


def test_coerce_excel_data_value_preserves_boolean_cells():
    """Booleans are not mistaken for integer values."""
    assert coerce_excel_data_value(True) is True


def test_write_data_keeps_workbook_metadata_as_text(tmp_path, monkeypatch):
    """Only table data cells are coerced; metadata and headers stay as text."""
    monkeypatch.setattr(excel_writers, "OUTPUT_LOCATION", tmp_path)
    excel_writers.create_empty_workbook("test", ["Sheet1"])

    excel_writers.write_data(
        pd.DataFrame({2025: ["2023"], "Value": ["0.95"], "Name": ["ELC_*"]}),
        book_name="test",
        sheet_name="Sheet1",
        tag=123,
        uc_set={},
        table_name=456,
        table_description=789,
    )

    sheet = load_workbook(tmp_path / "test.xlsx")["Sheet1"]
    assert sheet.cell(2, 1).value == "123"
    assert sheet.cell(2, 1).data_type == "s"
    assert sheet.cell(3, 1).value == "2025"
    assert sheet.cell(3, 1).data_type == "s"
    assert sheet.cell(4, 1).value == 2023
    assert sheet.cell(4, 2).value == 0.95
    assert sheet.cell(4, 3).value == "ELC_*"


def test_write_data_can_treat_a_single_value_header_as_data(tmp_path, monkeypatch):
    """The StartYear-style layout can retain a numeric header cell."""
    monkeypatch.setattr(excel_writers, "OUTPUT_LOCATION", tmp_path)
    excel_writers.create_empty_workbook("test", ["Sheet1"])

    excel_writers.write_data(
        pd.DataFrame(columns=[2023]),
        book_name="test",
        sheet_name="Sheet1",
        tag="~StartYear",
        uc_set={},
        header_is_data=True,
    )

    header_cell = load_workbook(tmp_path / "test.xlsx")["Sheet1"].cell(3, 1)
    assert header_cell.value == 2023
    assert header_cell.data_type == "n"
