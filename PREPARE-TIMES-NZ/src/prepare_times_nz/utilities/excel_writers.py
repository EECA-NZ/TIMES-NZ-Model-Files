"""
# bit of a mess will need to come back to this one

# contain all functions to read the metadata file in data_intermediate/config
# and use this to:

# 1: create all the workbooks and their relevant sheets
# 2: Identify all the tags we're going to write and their metadata
# (book, sheet, tag, uc_sets, maybe more)
# 3: get all the dataframes from the original toml files



"""

# Libraries ------------------------------------------------------------------
import os
import re
from decimal import ROUND_HALF_EVEN, Decimal, InvalidOperation
from numbers import Integral, Real

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from prepare_times_nz.utilities.filepaths import OUTPUT_LOCATION
from prepare_times_nz.utilities.logger_setup import logger

INTEGER_TEXT = re.compile(r"^[+-]?\d+$")
FLOAT_TEXT = re.compile(
    r"^[+-]?(?:(?:\d+\.\d*|\.\d+)(?:[eE][+-]?\d+)?|\d+[eE][+-]?\d+)$"
)


# Functions ------------------------------------------------------------------
def _round_decimal_to_significant_figures(
    value: Decimal, significant_figures: int = 15
) -> Decimal:
    """Round a finite Decimal to the precision Excel can retain."""
    if value.is_zero():
        return value

    exponent = value.adjusted() - significant_figures + 1
    return value.quantize(Decimal("1").scaleb(exponent), rounding=ROUND_HALF_EVEN)


def _excel_float(value: str | float | Decimal) -> float | str:
    """Return a 15-significant-figure float, or preserve non-finite text."""
    try:
        decimal_value = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return str(value)

    if not decimal_value.is_finite():
        return str(value)

    return float(_round_decimal_to_significant_figures(decimal_value))


def _coerce_excel_text(value: str) -> int | float | str:
    """Convert numeric-form text while preserving all other text verbatim."""
    if INTEGER_TEXT.fullmatch(value):
        return int(value)
    if FLOAT_TEXT.fullmatch(value):
        return _excel_float(value)
    return value


def coerce_excel_data_value(value):
    """Convert a table data value to its intended Excel cell type.

    CSV inputs are read as text so their spelling can be inspected before
    conversion. Integer-form text becomes an ``int``; decimal or scientific
    notation becomes a ``float`` rounded to 15 significant figures, Excel's
    usable numeric precision. All other values remain text. This function is
    deliberately applied only to table data cells: headers, tags, and other
    workbook metadata are always written as strings.
    """
    if value is None or pd.isna(value):
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, Integral):
        return int(value)

    if isinstance(value, (Real, Decimal)):
        return _excel_float(value)

    if isinstance(value, str):
        return _coerce_excel_text(value)

    return str(value)


def get_csv_data(file_location):
    """
    Returns the datafrom from a file_location,
    specifically we are ensuring that all data is read as strings
    This ensures full precision
    """
    df = pd.read_csv(file_location, dtype=str)
    return df


def dict_to_dataframe(data_dict):
    """
    takes a single dictionary from our tomls and creates a dataframe
    this is only used for the direct toml data
    (ie: dataframes entered into the toml files directly are converted here)
    """

    df_parts = []
    for key, values in data_dict.items():
        # Create DataFrame with explicit index
        if isinstance(values, (str, int, float, bool)) or values is None:
            # For scalar values, we need to wrap in a list and provide an index
            df_part = pd.DataFrame({key: [values]})
        else:
            # For dictionary or other iterable values
            df_part = pd.DataFrame({key: values})
        df_parts.append(df_part)

    # Concatenate all DataFrames
    if df_parts:
        df = pd.concat(df_parts, axis=1)
        df = df.fillna("")
    else:
        # Handle empty dictionary case
        df = pd.DataFrame()

    return df


def strip_headers_from_tiny_df(df):
    """
    We need special handling for some tables with one value
    XL2TIMES outputs these with a header called VALUE,
    but we need the value to be in the header column (with no data underneath)
    This function replaces the header with the value for specific tables

    """
    df = df.T  # Transpose the df
    df.columns = [df.iloc[0][0]]  # Set the column name to the value
    df = df.iloc[0:0]  # Remove all rows
    return df


def test_if_toml_location(string):
    """
    Return True if filename ends with .toml else False
    Mostly just wrapping the function in a more obvious name for its purpose

    """
    return string.endswith(".toml")


def create_empty_workbook(book_name, sheets):
    """
    Creates a workbook with empty sheets
    THe sheets will have data appended via excel overlay,
    but the sheets are required to exist first for that to work.

    Takes a book name and list of sheets, then saves the empty book
    in OUTPUT_LOCATION

    """
    # This function creates the workbook with empty sheets
    # Later, data is appended to these sheets by overlay.
    book_location = f"{OUTPUT_LOCATION}/{book_name}.xlsx"
    book_directory = os.path.dirname(book_location)

    # create the folder if needed
    os.makedirs(book_directory, exist_ok=True)

    # create workbook:
    wb = Workbook()

    # remove the default sheet which gets activated
    wb.remove(wb.active)

    # add each sheet
    for sheet in sheets:
        wb.create_sheet(sheet)

    # save
    wb.save(book_location)


def write_data(
    df,
    book_name,
    sheet_name,
    tag,
    uc_set,
    startrow=0,
    table_name="TABLE NAME",
    table_description="TABLE DESCRIPTION",
    header_is_data=False,
):
    """
    Writes data to the existing book

    Based on the input dataframe, the excel workbook is updated based on:

      The book's name
      THe Sheet name,
      The required input tag,
      Any uc_sets if relevant
      THe start row to begin adding data to

    The startrow input is a required input because each new tag added to a sheet
    needs to increase the start row so our tables don't overwrite each other.

    THis allows us to print many tagged tables to a single worksheet

    ``header_is_data`` is reserved for VEDA's unusual one-value-table layout,
    where the value is placed in a column header rather than a data row.

    """

    # Ideally, this function should be broken down into smaller pieces
    # to reduce the required arguments

    # pylint: disable=too-many-positional-arguments
    # pylint: disable=too-many-arguments
    # pylint: disable=too-many-locals

    # Load existing workbook
    book = load_workbook(f"{OUTPUT_LOCATION}/{book_name}.xlsx")
    sheet = book[sheet_name]
    # expand the width of the first column (often important for names)
    sheet.column_dimensions["A"].width = 30

    # Quick definitions of row layout. UC sets share the row immediately
    # above the table tag so xl2times associates them with that table without
    # treating a neighbouring cell as the tag's scalar value.
    description_row = startrow + 1
    uc_set_length = len(uc_set)
    if uc_set_length > 0:
        logger.debug("uc_sets detected")
    else:
        logger.debug("no uc_sets detected ")

    # A UC-set row is inserted between the description and table tag.
    tag_row = startrow + 2 + bool(uc_set_length)
    header_row = tag_row + 1
    data_start_row = tag_row + 2

    # Now can write things:
    # Write the table name and description, inherited from config files
    sheet.cell(row=description_row, column=1, value=str(table_name))
    sheet.cell(row=description_row, column=2, value=str(table_description))
    # make bold, for fun
    sheet.cell(row=description_row, column=1).font = Font(bold=True)
    sheet.cell(row=description_row, column=2).font = Font(bold=True)

    # Write the header row
    for col_idx, column_name in enumerate(df.columns, 1):
        header_value = (
            coerce_excel_data_value(column_name) if header_is_data else str(column_name)
        )
        sheet.cell(row=header_row, column=col_idx, value=header_value)

    # Write the data
    for row_idx, row in enumerate(df.values, data_start_row):
        for col_idx, value in enumerate(row, 1):
            sheet.cell(
                row=row_idx,
                column=col_idx,
                value=coerce_excel_data_value(value),
            )

    # Write the tag
    sheet.cell(row=tag_row, column=1, value=str(tag))

    # Add UC_Set tags if needed
    if uc_set_length > 0:
        uc_set_tag_row = tag_row - 1
        for column, (key, value) in enumerate(uc_set.items(), start=1):
            sheet.cell(
                row=uc_set_tag_row,
                column=column,
                value=f"~UC_Sets: {key}: {value}",
            )

    # Save the workbook
    book.save(f"{OUTPUT_LOCATION}/{book_name}.xlsx")
